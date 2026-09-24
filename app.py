"""
اپلیکیشن استخراج تردد ناقص و تردد چندگانه
شرکت ارس تارلا امین — یک فایل واحد Streamlit
"""

from __future__ import annotations

import io
import re
from collections import defaultdict
from datetime import time
from typing import Any, Optional

import openpyxl
import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# ثابت‌ها و ابزارها
# ---------------------------------------------------------------------------

PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")

COL_STATUS = 1   # A وضعیت
COL_GATE_OUT = 4  # D گیت خروج
COL_EXIT = 5      # E خروج
COL_GATE_IN = 7   # G گیت ورود
COL_ENTRY = 10    # J ورود
COL_SHIFT_END = 11    # K پایان شیفت
COL_SHIFT_START = 12  # L شروع شیفت
COL_DAY_TYPE = 14     # N نوع روز
COL_DAY = 15          # O روز
COL_DATE = 18         # R تاریخ
COL_NAME = 19         # S نام و نام خانوادگی
COL_CODE = 20         # T کدپرسنلی
COL_UNIT = 21         # U واحد سازمانی


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    for ch in ("\u200f", "\u200e", "\u200c", "\xa0"):
        s = s.replace(ch, " " if ch == "\u200c" else "")
    s = s.translate(PERSIAN_DIGITS).strip()
    return re.sub(r"\s+", " ", s)


def parse_time(value: Any) -> Optional[time]:
    s = clean_text(value)
    if not s or s in ("00:00", "0:00", "0:0"):
        return None
    m = re.match(r"^(\d{1,2}):(\d{2})$", s)
    if not m:
        return None
    h, mi = int(m.group(1)), int(m.group(2))
    if h > 23 or mi > 59:
        return None
    return time(h, mi)


def fmt_time(t: Optional[time]) -> str:
    return t.strftime("%H:%M") if t else ""


def time_to_minutes(t: time) -> int:
    return t.hour * 60 + t.minute


def classify_punch(
    punch: time,
    shift_start: Optional[time],
    shift_end: Optional[time],
    original: str,
) -> str:
    """طبقه‌بندی ورود/خروج بر اساس نزدیکی به شروع و پایان شیفت."""
    if shift_start is None or shift_end is None:
        return original
    p = time_to_minutes(punch)
    s = time_to_minutes(shift_start)
    e = time_to_minutes(shift_end)
    # شیفت شبانه
    if e <= s:
        e += 24 * 60
        if p < s:
            p += 24 * 60
    d_start = abs(p - s)
    d_end = abs(p - e)
    if d_start < d_end:
        return "in"
    if d_end < d_start:
        return "out"
    return original


def is_total_label(text: str) -> bool:
    t = text.replace("ـ", "").replace("-", "").replace(" ", "")
    return "مجموع" in t


def is_holiday(day_type: str) -> bool:
    return "تعط" in day_type


def is_leave_status(status: str) -> bool:
    return "مرخص" in status


def find_header_row(ws) -> int:
    for r in range(1, min(40, ws.max_row + 1)):
        for c in range(1, ws.max_column + 1):
            if clean_text(ws.cell(r, c).value) == "کدپرسنلی":
                return r
    raise ValueError("سطر هدر (کدپرسنلی) در فایل پیدا نشد.")


# ---------------------------------------------------------------------------
# پردازش اصلی
# ---------------------------------------------------------------------------

def process_workbook(file_bytes: bytes) -> dict[str, Any]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    header_row = find_header_row(ws)
    data_start = header_row + 2  # زیرسطر زیرستون‌ها

    cleaned_rows: list[dict] = []
    groups: dict[tuple, dict] = defaultdict(
        lambda: {"meta": {}, "punches": []}
    )

    cur_code = cur_name = cur_unit = None
    cur_date = cur_day = cur_dtype = None
    cur_start: Optional[time] = None
    cur_end: Optional[time] = None

    for r in range(data_start, ws.max_row + 1):
        status = clean_text(ws.cell(r, COL_STATUS).value)
        gate_out = clean_text(ws.cell(r, COL_GATE_OUT).value)
        t_out = parse_time(ws.cell(r, COL_EXIT).value)
        gate_in = clean_text(ws.cell(r, COL_GATE_IN).value)
        t_in = parse_time(ws.cell(r, COL_ENTRY).value)
        end_s = parse_time(ws.cell(r, COL_SHIFT_END).value)
        start_s = parse_time(ws.cell(r, COL_SHIFT_START).value)
        dtype = clean_text(ws.cell(r, COL_DAY_TYPE).value)
        day = clean_text(ws.cell(r, COL_DAY).value)
        date = clean_text(ws.cell(r, COL_DATE).value)
        name = clean_text(ws.cell(r, COL_NAME).value)
        code = clean_text(ws.cell(r, COL_CODE).value)
        unit = clean_text(ws.cell(r, COL_UNIT).value)

        if is_total_label(date) or is_total_label(name):
            continue
        if not any([status, gate_out, t_out, gate_in, t_in, date, name, code]):
            continue

        if code:
            cur_code = code
        if name:
            cur_name = name
        if unit:
            cur_unit = unit

        if date:
            cur_date = date
            cur_day = day
            # نوع روز را برای هر تاریخ جدید ریست می‌کنیم تا تعطیل ارث نبرد
            cur_dtype = dtype
            if start_s:
                cur_start = start_s
            if end_s:
                cur_end = end_s
        else:
            # ردیف ادامه‌دار همان تاریخ
            if day:
                cur_day = day
            if dtype:
                cur_dtype = dtype
            if start_s:
                cur_start = start_s
            if end_s:
                cur_end = end_s

        if not cur_code or not cur_date:
            continue

        # ردیف پاک‌سازی‌شده (همه ستون‌ها در یک سطح)
        cleaned_rows.append(
            {
                "کدپرسنلی": cur_code,
                "نام و نام خانوادگی": cur_name or "",
                "واحد سازمانی": cur_unit or "",
                "تاریخ": cur_date,
                "روز": cur_day or "",
                "نوع روز": cur_dtype or "",
                "شروع شیفت": fmt_time(cur_start),
                "پایان شیفت": fmt_time(cur_end),
                "ورود": fmt_time(t_in),
                "گیت ورود": gate_in,
                "خروج": fmt_time(t_out),
                "گیت خروج": gate_out,
                "وضعیت": status,
            }
        )

        if is_holiday(cur_dtype or ""):
            continue

        key = (cur_code, cur_date)
        groups[key]["meta"] = {
            "code": cur_code,
            "name": cur_name or "",
            "unit": cur_unit or "",
            "date": cur_date,
            "day": cur_day or "",
            "dtype": cur_dtype or "",
            "start": cur_start,
            "end": cur_end,
        }

        # رکورد مرخصی: ساعت‌های روی همان ردیف را نادیده می‌گیریم
        if is_leave_status(status):
            continue

        if t_in:
            kind = classify_punch(t_in, cur_start, cur_end, "in")
            groups[key]["punches"].append(
                {"time": t_in, "gate": gate_in, "kind": kind, "orig": "in"}
            )
        if t_out:
            kind = classify_punch(t_out, cur_start, cur_end, "out")
            groups[key]["punches"].append(
                {"time": t_out, "gate": gate_out, "kind": kind, "orig": "out"}
            )

    incomplete: list[dict] = []
    multiple: list[dict] = []

    for _key, g in groups.items():
        punches = g["punches"]
        meta = g["meta"]
        if not punches:
            continue

        ins = [p for p in punches if p["kind"] == "in"]
        outs = [p for p in punches if p["kind"] == "out"]

        base = {
            "کدپرسنلی": meta["code"],
            "نام و نام خانوادگی": meta["name"],
            "واحد سازمانی": meta["unit"],
            "تاریخ": meta["date"],
            "روز": meta["day"],
            "شروع شیفت": fmt_time(meta["start"]),
            "پایان شیفت": fmt_time(meta["end"]),
        }

        if len(ins) > 1 or len(outs) > 1:
            multiple.append(
                {
                    **base,
                    "تعداد ورود": len(ins),
                    "تعداد خروج": len(outs),
                    "زمان‌های ورود": "، ".join(
                        f"{fmt_time(p['time'])}"
                        + (f" ({p['gate']})" if p["gate"] else "")
                        for p in ins
                    ),
                    "زمان‌های خروج": "، ".join(
                        f"{fmt_time(p['time'])}"
                        + (f" ({p['gate']})" if p["gate"] else "")
                        for p in outs
                    ),
                    "شرح وضعیت": f"{len(ins)} ورود و {len(outs)} خروج",
                }
            )
        elif len(ins) == 0 or len(outs) == 0:
            incomplete.append(
                {
                    **base,
                    "ورود": fmt_time(ins[0]["time"]) if ins else "",
                    "گیت ورود": ins[0]["gate"] if ins else "",
                    "خروج": fmt_time(outs[0]["time"]) if outs else "",
                    "گیت خروج": outs[0]["gate"] if outs else "",
                    "نقص": "عدم ثبت ورود" if not ins else "عدم ثبت خروج",
                }
            )

    # مرتب‌سازی
    incomplete.sort(key=lambda x: (x["کدپرسنلی"], x["تاریخ"]))
    multiple.sort(key=lambda x: (x["کدپرسنلی"], x["تاریخ"]))

    cleaned_df = pd.DataFrame(cleaned_rows)
    incomplete_df = pd.DataFrame(incomplete) if incomplete else pd.DataFrame(
        columns=[
            "کدپرسنلی", "نام و نام خانوادگی", "واحد سازمانی", "تاریخ", "روز",
            "شروع شیفت", "پایان شیفت", "ورود", "گیت ورود", "خروج", "گیت خروج", "نقص",
        ]
    )
    multiple_df = pd.DataFrame(multiple) if multiple else pd.DataFrame(
        columns=[
            "کدپرسنلی", "نام و نام خانوادگی", "واحد سازمانی", "تاریخ", "روز",
            "شروع شیفت", "پایان شیفت", "تعداد ورود", "تعداد خروج",
            "زمان‌های ورود", "زمان‌های خروج", "شرح وضعیت",
        ]
    )

    # خروجی نهایی تردد ناقص با ستون‌های درخواستی کاربر
    incomplete_export = incomplete_df[
        [
            "کدپرسنلی",
            "نام و نام خانوادگی",
            "تاریخ",
            "روز",
            "ورود",
            "گیت ورود",
            "خروج",
            "گیت خروج",
        ]
    ].copy() if not incomplete_df.empty else incomplete_df[
        ["کدپرسنلی", "نام و نام خانوادگی", "تاریخ", "روز", "ورود", "گیت ورود", "خروج", "گیت خروج"]
    ]
    if not incomplete_df.empty:
        incomplete_export["نقص"] = incomplete_df["نقص"]

    return {
        "cleaned": cleaned_df,
        "incomplete": incomplete_df,
        "incomplete_export": incomplete_export,
        "multiple": multiple_df,
    }


def style_worksheet(ws, header_fill_hex: str = "1F4D3A") -> None:
    header_fill = PatternFill("solid", fgColor=header_fill_hex)
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    body_font = Font(name="Arial", size=10)
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.font = body_font
            cell.border = thin
            cell.alignment = center if cell.column <= 4 else right

    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for cell in ws[col_letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 40)

    ws.row_dimensions[1].height = 28
    ws.sheet_view.rightToLeft = True


def df_to_excel_bytes(sheets: dict[str, pd.DataFrame], colors: dict[str, str] | None = None) -> bytes:
    colors = colors or {}
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
            ws = writer.sheets[name[:31]]
            style_worksheet(ws, colors.get(name, "1F4D3A"))
    buf.seek(0)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# رابط کاربری Streamlit
# ---------------------------------------------------------------------------

st.set_page_config(
    page_title="استخراج تردد ناقص | ارس تارلا امین",
    page_icon="📋",
    layout="wide",
    initial_sidebar_state="expanded",
)

CUSTOM_CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Vazirmatn:wght@400;500;600;700&display=swap');

html, body, [class*="css"] {
    font-family: 'Vazirmatn', Tahoma, sans-serif !important;
    direction: rtl;
    text-align: right;
}

.stApp {
    background: linear-gradient(160deg, #f4efe6 0%, #e8f0eb 45%, #f7f3ec 100%);
}

h1, h2, h3, h4 {
    font-family: 'Vazirmatn', Tahoma, sans-serif !important;
    color: #1f4d3a !important;
}

div[data-testid="stSidebar"] {
    background: #1f4d3a;
}
div[data-testid="stSidebar"] * {
    color: #f4efe6 !important;
}
div[data-testid="stSidebar"] .stMarkdown p {
    color: #d8e8df !important;
}

.hero {
    background: #1f4d3a;
    color: #f4efe6;
    padding: 1.4rem 1.6rem;
    border-radius: 16px;
    margin-bottom: 1.2rem;
    box-shadow: 0 8px 24px rgba(31, 77, 58, 0.18);
}
.hero h1 {
    color: #f4efe6 !important;
    margin: 0 0 0.35rem 0;
    font-size: 1.55rem;
    font-weight: 700;
}
.hero p {
    margin: 0;
    opacity: 0.9;
    font-size: 0.95rem;
}

.metric-card {
    background: #fffaf3;
    border: 1px solid #e2d6c5;
    border-radius: 14px;
    padding: 1rem 1.1rem;
    text-align: center;
    box-shadow: 0 2px 8px rgba(0,0,0,0.04);
}
.metric-card .num {
    font-size: 1.8rem;
    font-weight: 700;
    color: #1f4d3a;
    line-height: 1.2;
}
.metric-card .label {
    font-size: 0.85rem;
    color: #6b6459;
    margin-top: 0.25rem;
}
.metric-card.warn .num { color: #c45c26; }
.metric-card.multi .num { color: #8b3a62; }

.stDownloadButton button, .stButton button {
    font-family: 'Vazirmatn', Tahoma, sans-serif !important;
    border-radius: 10px !important;
}

div[data-testid="stFileUploader"] {
    background: #fffaf3;
    border: 2px dashed #1f4d3a;
    border-radius: 14px;
    padding: 0.8rem;
}

.hint {
    background: #fffaf3;
    border-right: 4px solid #c45c26;
    padding: 0.85rem 1rem;
    border-radius: 8px;
    color: #3a342c;
    margin: 0.8rem 0 1.2rem 0;
    font-size: 0.92rem;
}

footer { visibility: hidden; }
</style>
"""
st.markdown(CUSTOM_CSS, unsafe_allow_html=True)

st.markdown(
    """
    <div class="hero">
      <h1>استخراج تردد ناقص و تردد چندگانه</h1>
      <p>شرکت ارس تارلا امین (سهامی خاص) — پاک‌سازی گزارش تردد روزانه و شناسایی ثبت‌های ناقص یا تکراری</p>
    </div>
    """,
    unsafe_allow_html=True,
)

with st.sidebar:
    st.markdown("### راهنما")
    st.markdown(
        """
        ۱. فایل اکسل گزارش تردد را آپلود کنید  
        ۲. سیستم ستون‌های «اطلاعات تردد» را تخت می‌کند  
        ۳. با مقایسه ساعت تردد و شروع/پایان شیفت، ورود و خروج واقعی تشخیص داده می‌شود  
        ۴. دو خروجی ساخته می‌شود:
        - **تردد ناقص**: فقط یک سمت (ورود یا خروج) ثبت شده
        - **تردد چندگانه**: بیش از یک ورود یا خروج در یک تاریخ
        """
    )
    st.markdown("---")
    st.markdown("### جستجوی پرسنل")
    search_code = st.text_input(
        "کد پرسنلی",
        placeholder="مثال: 110864",
        help="پس از پردازش، نتایج مربوط به این کد نمایش داده می‌شود.",
    )
    search_name = st.text_input(
        "نام (اختیاری)",
        placeholder="بخشی از نام...",
    )

uploaded = st.file_uploader(
    "فایل اکسل گزارش تردد را انتخاب کنید",
    type=["xlsx", "xls"],
    help="خروجی سیستم حضور و غیاب با ساختار استاندارد شرکت",
)

# بارگذاری نمونه برای دمو در صورت نبود فایل
SAMPLE_PATH = "/home/workdir/attachments/تردد-ناقص-اشخاص.xlsx"
use_sample = False
if uploaded is None:
    st.markdown(
        '<div class="hint">فایلی آپلود نشده است. می‌توانید فایل نمونه ضمیمه‌شده را برای تست پردازش کنید.</div>',
        unsafe_allow_html=True,
    )
    use_sample = st.checkbox("استفاده از فایل نمونه ضمیمه‌شده", value=True)

if uploaded is not None or use_sample:
    try:
        if uploaded is not None:
            file_bytes = uploaded.getvalue()
            source_label = uploaded.name
        else:
            with open(SAMPLE_PATH, "rb") as f:
                file_bytes = f.read()
            source_label = "تردد-ناقص-اشخاص.xlsx (نمونه)"

        with st.spinner("در حال پاک‌سازی و استخراج..."):
            result = process_workbook(file_bytes)

        cleaned = result["cleaned"]
        incomplete = result["incomplete"]
        incomplete_export = result["incomplete_export"]
        multiple = result["multiple"]

        c1, c2, c3, c4 = st.columns(4)
        with c1:
            st.markdown(
                f'<div class="metric-card"><div class="num">{len(cleaned)}</div>'
                f'<div class="label">ردیف پاک‌سازی‌شده</div></div>',
                unsafe_allow_html=True,
            )
        with c2:
            st.markdown(
                f'<div class="metric-card warn"><div class="num">{len(incomplete)}</div>'
                f'<div class="label">تردد ناقص</div></div>',
                unsafe_allow_html=True,
            )
        with c3:
            st.markdown(
                f'<div class="metric-card multi"><div class="num">{len(multiple)}</div>'
                f'<div class="label">تردد چندگانه</div></div>',
                unsafe_allow_html=True,
            )
        with c4:
            people = set(incomplete["کدپرسنلی"].tolist()) | set(multiple["کدپرسنلی"].tolist()) if (
                not incomplete.empty or not multiple.empty
            ) else set()
            st.markdown(
                f'<div class="metric-card"><div class="num">{len(people)}</div>'
                f'<div class="label">تعداد افراد درگیر</div></div>',
                unsafe_allow_html=True,
            )

        st.caption(f"منبع: {source_label}")

        # فیلتر جستجو
        code_q = clean_text(search_code)
        name_q = clean_text(search_name)

        def apply_filter(df: pd.DataFrame) -> pd.DataFrame:
            if df.empty:
                return df
            out = df
            if code_q:
                out = out[out["کدپرسنلی"].astype(str).str.contains(code_q, na=False)]
            if name_q:
                out = out[out["نام و نام خانوادگی"].astype(str).str.contains(name_q, na=False)]
            return out

        filt_incomplete = apply_filter(incomplete)
        filt_multiple = apply_filter(multiple)

        tab1, tab2, tab3 = st.tabs(
            ["تردد ناقص", "تردد چندگانه", "داده پاک‌سازی‌شده"]
        )

        with tab1:
            st.markdown(
                "روزهایی که **فقط ورود یا فقط خروج** ثبت شده "
                "(با اصلاح طبقه‌بندی بر اساس شروع/پایان شیفت)."
            )
            if filt_incomplete.empty:
                st.success("موردی یافت نشد.")
            else:
                display_cols = [
                    "کدپرسنلی", "نام و نام خانوادگی", "تاریخ", "روز",
                    "ورود", "گیت ورود", "خروج", "گیت خروج", "نقص",
                    "شروع شیفت", "پایان شیفت",
                ]
                st.dataframe(
                    filt_incomplete[display_cols],
                    use_container_width=True,
                    hide_index=True,
                )

            xls_incomplete = df_to_excel_bytes(
                {"تردد ناقص": incomplete_export if not incomplete_export.empty else incomplete},
                {"تردد ناقص": "C45C26"},
            )
            st.download_button(
                label="دانلود اکسل تردد ناقص",
                data=xls_incomplete,
                file_name="تردد_ناقص.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
            )

        with tab2:
            st.markdown(
                "روزهایی که در یک تاریخ خاص **بیش از یک ورود یا بیش از یک خروج** ثبت شده است."
            )
            if filt_multiple.empty:
                st.success("موردی یافت نشد.")
            else:
                st.dataframe(
                    filt_multiple,
                    use_container_width=True,
                    hide_index=True,
                )

            xls_multi = df_to_excel_bytes(
                {"تردد چندگانه": multiple},
                {"تردد چندگانه": "8B3A62"},
            )
            st.download_button(
                label="دانلود اکسل تردد چندگانه",
                data=xls_multi,
                file_name="تردد_چندگانه.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
            )

        with tab3:
            st.markdown(
                "ستون «اطلاعات تردد» حذف و زیرستون‌ها (وضعیت، ورود، گیت، خروج، گیت) "
                "هم‌ردیف با شروع/پایان شیفت قرار گرفته‌اند. ردیف‌های مجموع حذف شده‌اند."
            )
            st.dataframe(cleaned, use_container_width=True, hide_index=True)
            xls_clean = df_to_excel_bytes(
                {"داده پاک‌سازی‌شده": cleaned},
                {"داده پاک‌سازی‌شده": "1F4D3A"},
            )
            st.download_button(
                label="دانلود اکسل داده پاک‌سازی‌شده",
                data=xls_clean,
                file_name="داده_پاکسازی_شده.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        # دانلود یکجا
        st.markdown("---")
        all_bytes = df_to_excel_bytes(
            {
                "تردد ناقص": incomplete_export if not incomplete_export.empty else incomplete,
                "تردد چندگانه": multiple,
                "داده پاک‌سازی‌شده": cleaned,
            },
            {
                "تردد ناقص": "C45C26",
                "تردد چندگانه": "8B3A62",
                "داده پاک‌سازی‌شده": "1F4D3A",
            },
        )
        st.download_button(
            label="دانلود همه خروجی‌ها در یک فایل",
            data=all_bytes,
            file_name="گزارش_تردد_کامل.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        # نمایش نمونه استثناها در صورت جستجوی کد شناخته‌شده
        if code_q in ("110864", "112332"):
            st.info(
                "این کد یکی از نمونه‌های استثنایی فایل است: "
                "طبقه‌بندی ورود/خروج با کنترل شروع و پایان شیفت انجام شده "
                "و ترددهای چندگانه در برگه جداگانه قرار گرفته‌اند."
            )

    except Exception as exc:
        st.error(f"خطا در پردازش فایل: {exc}")
        st.exception(exc)
else:
    st.info("برای شروع، یک فایل اکسل گزارش تردد آپلود کنید.")
