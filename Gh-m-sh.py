import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill
import io
import re

# تابع نرمال‌سازی کدهای پرسنلی (تبدیل اعداد فارسی به انگلیسی و حذف صفرهای قبل از عدد)
def normalize_code(text):
    if pd.isna(text) or text is None or text == 'None':
        return ""
    text = str(text).strip()
    persian_to_english = str.maketrans('۰۱۲۳۴۵۶۷۸۹', '0123456789')
    arabic_to_english = str.maketrans('٠١٢٣٤٥٦٧٨٩', '0123456789')
    text = text.translate(persian_to_english).translate(arabic_to_english)
    text = text.lstrip('0')
    if '.' in text:
        text = text.split('.')[0]
    return text

# تابع استخراج تاریخ از متن
def extract_date(text):
    if pd.isna(text) or text is None or text == 'None':
        return ""
    text = str(text)
    persian_to_english = str.maketrans('۰۱۲۳۴۵۶۷۸۹', '0123456789')
    arabic_to_english = str.maketrans('٠١٢٣٤٥٦٧٨٩', '0123456789')
    text = text.translate(persian_to_english).translate(arabic_to_english)
    match = re.search(r'(\d{4}/\d{2}/\d{2})', text)
    if match:
        return match.group(1)
    return ""

st.set_page_config(page_title="بررسی غیبت و مرخصی", layout="wide")
st.title("تطبیق هوشمند غیبت‌ها با درخواست‌های مرخصی")
st.markdown("فایل **غیبت** و فایل **درخواست مرخصی** را آپلود کنید تا سیستم غیبت‌های موجه (دارای مرخصی) را تشخیص داده و خروجی‌ها را تولید کند.")

col1, col2 = st.columns(2)
with col1:
    file_abs = st.file_uploader("۱. فایل غیبت را آپلود کنید (xlsx)", type=['xlsx'])
with col2:
    file_leave = st.file_uploader("۲. فایل درخواست مرخصی را آپلود کنید (xlsx)", type=['xlsx'])

if st.button("پردازش و تطبیق", type="primary", use_container_width=True):
    if file_abs and file_leave:
        try:
            with st.spinner("در حال پردازش داده‌ها... این عملیات ممکن است چند ثانیه طول بکشد."):
                
                # --- بخش اول: پردازش فایل مرخصی ---
                df_leave = pd.read_excel(file_leave)
                
                leave_code_col = [c for c in df_leave.columns if 'کد پرسنلی' in str(c)][0]
                leave_start_col = [c for c in df_leave.columns if 'شروع' in str(c)][0]
                leave_end_col = [c for c in df_leave.columns if 'پایان' in str(c)][0]
                
                df_leave['کد_نرمال'] = df_leave[leave_code_col].apply(normalize_code)
                df_leave['شروع_نرمال'] = df_leave[leave_start_col].apply(extract_date)
                df_leave['پایان_نرمال'] = df_leave[leave_end_col].apply(extract_date)
                
                # --- بخش دوم: پردازش فایل غیبت ---
                wb = openpyxl.load_workbook(file_abs, data_only=False)
                ws = wb.active
                
                date_col_idx = None
                code_col_idx = None
                
                for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True)):
                    for col_idx, cell_value in enumerate(row):
                        val = str(cell_value).replace('\n', ' ').strip()
                        if val == "تاریخ":
                            date_col_idx = col_idx
                        if "کد پرسنلی" in val:
                            code_col_idx = col_idx
                
                if date_col_idx is None or code_col_idx is None:
                    st.error("ستون 'تاریخ' یا 'کد پرسنلی' در فایل غیبت یافت نشد.")
                    st.stop()

                # === رفع خطای ادغام (Merge) ===
                # خارج کردن سلول‌های ستون کد پرسنلی از حالت ادغام برای امکان مقداردهی
                for merged_range in list(ws.merged_cells.ranges):
                    if merged_range.min_col <= (code_col_idx + 1) <= merged_range.max_col:
                        ws.unmerge_cells(str(merged_range))
                
                current_code = ""
                rows_to_highlight = []
                yellow_fill = PatternFill(start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid')
                
                for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                    raw_code = row[code_col_idx]
                    norm_code = normalize_code(raw_code)
                    
                    if norm_code != "" and norm_code.isdigit():
                        current_code = norm_code
                    
                    # مقداردهی ردیف‌های خالی با کد پرسنلی (حالا بدون خطا انجام می‌شود)
                    if (raw_code is None or str(raw_code).strip() == "" or str(raw_code) == "None") and current_code != "":
                        ws.cell(row=row_idx, column=code_col_idx + 1).value = current_code
                    
                    raw_date = row[date_col_idx]
                    norm_date = extract_date(raw_date)
                    
                    if current_code != "" and norm_date != "":
                        personnel_leaves = df_leave[df_leave['کد_نرمال'] == current_code]
                        
                        is_justified = False
                        for _, leave_row in personnel_leaves.iterrows():
                            if leave_row['شروع_نرمال'] <= norm_date <= leave_row['پایان_نرمال']:
                                is_justified = True
                                break
                        
                        if is_justified:
                            rows_to_highlight.append(row_idx)
                
                # --- خروجی اول: هایلایت ---
                for r_idx in rows_to_highlight:
                    for cell in ws[r_idx]:
                        cell.fill = yellow_fill
                
                output_highlighted = io.BytesIO()
                wb.save(output_highlighted)
                output_highlighted.seek(0)
                
                # --- خروجی دوم: حذف موجه‌ها ---
                for r_idx in sorted(rows_to_highlight, reverse=True):
                    ws.delete_rows(r_idx)
                
                output_cleaned = io.BytesIO()
                wb.save(output_cleaned)
                output_cleaned.seek(0)
                
            st.success(f"✅ پردازش تکمیل شد! تعداد {len(rows_to_highlight)} ردیف به عنوان غیبت موجه (ثبت شده در مرخصی) تشخیص داده شد.")
            
            st.divider()
            st.subheader("دریافت خروجی‌ها")
            
            col_out1, col_out2 = st.columns(2)
            with col_out1:
                st.download_button(
                    label="📥 دانلود فایل با غیبت‌های موجه (هایلایت شده)",
                    data=output_highlighted,
                    file_name="Absence_Highlighted.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
            
            with col_out2:
                st.download_button(
                    label="📥 دانلود فایل غیبت‌های غیرموجه (حذف موجه‌ها)",
                    data=output_cleaned,
                    file_name="Absence_Unjustified_Only.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
                
        except Exception as e:
            st.error(f"خطایی در پردازش رخ داد: {e}")
    else:
        st.warning("لطفاً هر دو فایل را آپلود کنید.")
