import streamlit as st
import pandas as pd
import io
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

# تنظیمات صفحه
st.set_page_config(page_title="تحلیل‌گر زمان کارکرد", layout="centered", page_icon="📊")

def time_to_minutes(time_str):
    """تبدیل فرمت hhh:mm به کل دقایق"""
    if pd.isna(time_str) or str(time_str).strip() == '':
        return None
    try:
        time_str = str(time_str).strip()
        if ':' in time_str:
            parts = time_str.split(':')
            hours = int(parts[0])
            minutes = int(parts[1])
            # مدیریت اعداد منفی
            if hours < 0 or time_str.startswith('-'):
                return (abs(hours) * 60 + minutes) * -1
            return hours * 60 + minutes
        else:
            return int(time_str) * 60
    except Exception:
        return None

def minutes_to_time(total_minutes):
    """تبدیل کل دقایق به فرمت hhh:mm"""
    if pd.isna(total_minutes) or total_minutes is None:
        return "00:00"
    
    sign = "-" if total_minutes < 0 else ""
    total_minutes = abs(int(total_minutes))
    hours = total_minutes // 60
    minutes = total_minutes % 60
    return f"{sign}{hours:02d}:{minutes:02d}"

st.title("📊 تحلیل‌گر اضافه‌کار و کسرکار")
st.write("فایل اکسل خود را آپلود کنید تا مجموع و میانگین ساعات محاسبه شده و در انتهای فایل با رنگ مشخص اضافه گردد.")

uploaded_file = st.file_uploader("آپلود فایل اکسل", type=['xlsx'])

if uploaded_file is not None:
    # خواندن فایل اکسل
    try:
        df = pd.read_excel(uploaded_file)
        
        col_overtime = 'اضافه کار عادی ماهانه نهایی'
        col_deduction = 'کسرکار نهایی ماهانه'
        
        # بررسی وجود ستون‌های مورد نیاز در فایل
        if col_overtime not in df.columns or col_deduction not in df.columns:
            st.error(f"خطا: ستون‌های '{col_overtime}' و '{col_deduction}' در فایل اکسل یافت نشدند.")
        else:
            st.success("فایل با موفقیت بارگذاری و ستون‌ها شناسایی شدند!")
            
            # تبدیل مقادیر به دقیقه و حذف مقادیر خالی برای محاسبه دقیق‌تر
            overtime_mins = df[col_overtime].apply(time_to_minutes).dropna()
            deduction_mins = df[col_deduction].apply(time_to_minutes).dropna()
            
            # محاسبه مجموع و میانگین (به دقیقه)
            sum_overtime_mins = overtime_mins.sum() if not overtime_mins.empty else 0
            mean_overtime_mins = overtime_mins.mean() if not overtime_mins.empty else 0
            
            sum_deduction_mins = deduction_mins.sum() if not deduction_mins.empty else 0
            mean_deduction_mins = deduction_mins.mean() if not deduction_mins.empty else 0
            
            # نمایش نتایج در داشبورد اپلیکیشن
            st.markdown("### 📈 نتایج تحلیل:")
            col1, col2 = st.columns(2)
            
            with col1:
                st.info(f"**{col_overtime}**")
                st.write(f"مجموع: **{minutes_to_time(sum_overtime_mins)}**")
                st.write(f"میانگین: **{minutes_to_time(mean_overtime_mins)}**")
                
            with col2:
                st.warning(f"**{col_deduction}**")
                st.write(f"مجموع: **{minutes_to_time(sum_deduction_mins)}**")
                st.write(f"میانگین: **{minutes_to_time(mean_deduction_mins)}**")
            
            # آماده‌سازی ردیف‌های جدید برای افزودن به دیتافریم
            sum_row = pd.Series(index=df.columns, dtype=object)
            mean_row = pd.Series(index=df.columns, dtype=object)
            
            # قرار دادن عناوین در ستون اول (یا ستون مناسب دیگر)
            first_col = df.columns[0]
            sum_row[first_col] = "مجموع کل:"
            mean_row[first_col] = "میانگین کل:"
            
            # مقداردهی به ستون‌های هدف
            sum_row[col_overtime] = minutes_to_time(sum_overtime_mins)
            sum_row[col_deduction] = minutes_to_time(sum_deduction_mins)
            
            mean_row[col_overtime] = minutes_to_time(mean_overtime_mins)
            mean_row[col_deduction] = minutes_to_time(mean_deduction_mins)
            
            # اضافه کردن ردیف‌ها به انتهای فایل
            df_export = pd.concat([df, sum_row.to_frame().T, mean_row.to_frame().T], ignore_index=True)
            
            # ذخیره دیتافریم در حافظه موقت (برای اعمال استایل روی فایل اکسل)
            output = io.BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df_export.to_excel(writer, index=False, sheet_name='Sheet1')
            
            # باز کردن فایل با openpyxl برای استایل‌دهی (هایلایت کردن)
            output.seek(0)
            wb = load_workbook(output)
            ws = wb.active
            
            # پیدا کردن ایندکس ستون‌های هدف
            col_idx_overtime = df_export.columns.get_loc(col_overtime) + 1
            col_idx_deduction = df_export.columns.get_loc(col_deduction) + 1
            
            # تعریف استایل: پس‌زمینه زرد رنگ و فونت بولد
            highlight_fill = PatternFill(start_color="FFFFE0", end_color="FFFF00", fill_type="solid")
            bold_font = Font(bold=True)
            
            # دو ردیف آخر که مجموع و میانگین هستند را هایلایت می‌کنیم
            max_row = ws.max_row
            for row_num in range(max_row - 1, max_row + 1):
                # استایل برای ستون اضافه کار
                ws.cell(row=row_num, column=col_idx_overtime).fill = highlight_fill
                ws.cell(row=row_num, column=col_idx_overtime).font = bold_font
                # استایل برای ستون کسر کار
                ws.cell(row=row_num, column=col_idx_deduction).fill = highlight_fill
                ws.cell(row=row_num, column=col_idx_deduction).font = bold_font
                # بولد کردن برچسب (مجموع/میانگین)
                ws.cell(row=row_num, column=1).font = bold_font

            # ذخیره فایل استایل‌دهی شده
            final_output = io.BytesIO()
            wb.save(final_output)
            final_output.seek(0)
            
            st.divider()
            st.markdown("### 📥 دانلود خروجی نهایی")
            st.download_button(
                label="دانلود فایل اکسل ویرایش شده",
                data=final_output,
                file_name="Processed_Export.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

    except Exception as e:
        st.error(f"مشکلی در پردازش فایل به وجود آمد: {e}")
